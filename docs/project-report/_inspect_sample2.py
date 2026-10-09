# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

path = Path(r"c:\Users\Artemis\Downloads\نمونه-فایل-فرمت-تحقیق-دانشجویی.docx")
out = Path(r"d:\Personal\code 2\goldon hoshmand\docs\project-report\_sample_detail.txt")
doc = Document(str(path))
lines = []

def w(s=""):
    lines.append(s)

w("SECTIONS=%d" % len(doc.sections))
for i, sec in enumerate(doc.sections):
    w(f"section {i}: {round(sec.page_width.cm,2)}x{round(sec.page_height.cm,2)} cm")
    w(f"  margins R={round(sec.right_margin.cm,2)} L={round(sec.left_margin.cm,2)} T={round(sec.top_margin.cm,2)} B={round(sec.bottom_margin.cm,2)}")
    w(f"  header dist={round(sec.header_distance.cm,2)} footer dist={round(sec.footer_distance.cm,2)}")
    w(f"  different_first={sec.different_first_page_header_footer}")
    hdr = sec.header
    ftr = sec.footer
    w(f"  header paras: {[p.text for p in hdr.paragraphs]}")
    w(f"  footer paras: {[p.text for p in ftr.paragraphs]}")
    # page number fields?
    xml = sec._sectPr.xml
    w(f"  sectPr snippet has pgNumType={'pgNumType' in xml}")

w("\n--- STYLES ---")
for st in doc.styles:
    try:
        name = st.name
        font = st.font
        fname = font.name
        size = font.size.pt if font.size else None
        bold = font.bold
        if name in (
            "Normal", "Heading 1", "Heading 2", "Heading 3",
            "Title", "Subtitle", "Title Page", "Title Page Names",
            "toc 1", "toc 2", "toc 3", "شکل", "جدول", "متن",
            "شماره گذاري مراجع", "Normal Text",
        ) or (fname or size):
            if any(x in (name or "") for x in (
                "Heading", "Title", "Normal", "toc", "شکل", "جدول", "متن", "مرجع", "Subtitle"
            )):
                w(f"style {name!r} font={fname} size={size} bold={bold}")
    except Exception:
        pass

w("\n--- DEFAULT THEME / DOCUMENT FONT ---")
# styles.xml default
from zipfile import ZipFile
with ZipFile(path) as z:
    styles = z.read("word/styles.xml").decode("utf-8", errors="replace")
    # rFonts
    import re
    fonts_found = sorted(set(re.findall(r'w:ascii="([^"]+)"', styles)))
    cs_fonts = sorted(set(re.findall(r'w:cs="([^"]+)"', styles)))
    w("ascii fonts: " + ", ".join(fonts_found[:40]))
    w("cs fonts: " + ", ".join(cs_fonts[:40]))
    sizes = sorted(set(re.findall(r'w:sz w:val="(\d+)"', styles)))
    w("sz vals: " + ", ".join(sizes[:30]))
    # theme
    if "word/theme/theme1.xml" in z.namelist():
        theme = z.read("word/theme/theme1.xml").decode("utf-8", errors="replace")
        maj = re.findall(r'typeface="([^"]+)"', theme)
        w("theme typefaces: " + ", ".join(sorted(set(maj))[:20]))
    w("files: " + ", ".join(z.namelist()[:40]))

w("\n--- FULL PARAS ---")
for i, p in enumerate(doc.paragraphs):
    t = (p.text or "").replace("\t", " | ")
    xml = p._p.xml
    pb = " PAGEBREAK" if ('w:type="page"' in xml or "w:type='page'" in xml) else ""
    if not t.strip() and not pb:
        continue
    align = p.alignment
    w(f"[{i}] style={p.style.name!r} align={align}{pb}")
    w(f"    {t[:300]}")

w("\n--- TABLES DETAIL ---")
for ti, table in enumerate(doc.tables):
    w(f"table {ti} {len(table.rows)}x{len(table.columns)}")
    for row in table.rows:
        cells = [c.text.replace("\n", " / ") for c in row.cells]
        w("  " + " || ".join(cells))

out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", out, "lines", len(lines))
